"""5등급.xlsx → data/ipgyeol.json, data/conv.json, test-data/golden.json
사용: python scripts/extract.py [엑셀경로]
"""
import json, sys, os
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\user\Desktop\5등급.xlsx"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COLUMNS = ["지역", "대학명", "전형", "전형명", "계열", "모집단위", "모집인원", "경쟁률", "충원순위", "cut70"]


def _cell(v):
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def _row(values):
    return [_cell(v) for v in values[:10]]


def dump(rel, obj):
    path = os.path.join(ROOT, rel)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))
    print(f"{rel}: {os.path.getsize(path):,} bytes")


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)

    # 1) 입결
    ws = wb["2026학년도 입결"]
    it = ws.iter_rows(values_only=True)
    header = [_cell(v) for v in next(it)[:10]]
    assert header[0] == "지역" and header[9] == "70% cut", header
    rows = [_row(r) for r in it if r and r[0]]
    assert len(rows) == 14612, len(rows)
    assert all(isinstance(r[9], (int, float)) for r in rows), "cut70 must be numeric"
    dump("data/ipgyeol.json", {"source": "2026학년도 입결", "columns": COLUMNS, "rows": rows})

    # 2) 변환표 T2:U396 + 골든 결과 J3:S
    ws = wb["결과시트"]
    pairs, golden, inp, conv = [], [], None, None
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        r = list(r) + [None] * (21 - len(r))
        if i == 1:
            inp = r[2]          # C1
        elif i == 2:
            conv = r[2]         # C2
        if i >= 2 and r[19] is not None:
            pairs.append([float(r[19]), float(r[20])])
        if i >= 3 and r[9]:
            golden.append(_row(r[9:19]))
    assert len(pairs) == 395, len(pairs)
    assert all(pairs[k][0] < pairs[k + 1][0] for k in range(len(pairs) - 1)), "5등급 키 오름차순 아님"
    assert all(pairs[k][1] <= pairs[k + 1][1] for k in range(len(pairs) - 1)), "9등급 값 비감소 아님"
    assert pairs[0][0] == 1 and pairs[-1][0] == 5
    dump("data/conv.json", {"source": "결과시트 T2:U396", "pairs": pairs})

    assert inp == 1.1 and conv == 1.39 and len(golden) == 259, (inp, conv, len(golden))
    lo, hi = round(conv - 0.1, 3), round(conv + 0.1, 3)
    dump("test-data/golden.json", {"input": inp, "converted": conv, "lo": lo, "hi": hi,
                                   "columns": COLUMNS, "rows": golden})


if __name__ == "__main__":
    main()
